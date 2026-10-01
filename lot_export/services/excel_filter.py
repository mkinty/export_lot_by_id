"""
Filtrage d'un fichier audit : ne conserve, dans la feuille Audit, que les lignes
dont la valeur d'une colonne donnée (« ID erreur », « adresse »...) fait partie d'une liste.
Les autres feuilles sont copiées telles quelles.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

AUDIT_SHEET = "Audit"
ERROR_ID_HEADER = "ID erreur"
ADDRESS_HEADER = "adresse"


class ExcelFilterError(Exception):
    """Structure du fichier audit inattendue (feuille ou colonne absente)."""


@dataclass
class FilterResult:
    """Valeurs demandées trouvées (kept_ids) ou absentes (missing_ids) de la colonne filtrée."""
    kept_ids: List[str] = field(default_factory=list)
    missing_ids: List[str] = field(default_factory=list)


def _normalize(value: object) -> str:
    """Comparaison insensible à la casse, aux accents et aux espaces multiples."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().upper()


def _find_column(ws: Worksheet, header: str) -> int:
    for cell in ws[1]:
        if _normalize(cell.value) == _normalize(header):
            return cell.column
    raise ExcelFilterError(f"colonne « {header} » introuvable")


def filter_audit_rows(
        source: Path, destination: Path, values: Iterable[str], column: str = ERROR_ID_HEADER,
) -> FilterResult:
    """
    Écrit dans destination une copie de source où la feuille Audit ne contient plus
    que l'en-tête et les lignes dont la valeur de la colonne `column` est dans values.
    N'écrit rien si aucune ligne ne correspond.
    """
    # clé normalisée -> valeur telle que saisie (pour le compte rendu)
    wanted = {}
    for value in values:
        wanted.setdefault(_normalize(value), value)

    wb = load_workbook(source)
    if AUDIT_SHEET not in wb.sheetnames:
        raise ExcelFilterError(f"feuille « {AUDIT_SHEET} » introuvable")
    ws = wb[AUDIT_SHEET]
    id_col = _find_column(ws, column)
    last_col = get_column_letter(ws.max_column)
    max_row = ws.max_row

    found: set[str] = set()
    target = 2
    for row in range(2, max_row + 1):
        value = _normalize(ws.cell(row, id_col).value)
        if value not in wanted:
            continue
        found.add(value)
        if row != target:
            # translate=True réécrit les formules relatives (ex: J57 -> J3) pour la nouvelle ligne
            ws.move_range(f"A{row}:{last_col}{row}", rows=target - row, translate=True)
        target += 1

    result = FilterResult(
        kept_ids=[original for key, original in wanted.items() if key in found],
        missing_ids=[original for key, original in wanted.items() if key not in found],
    )
    if not found:
        return result

    if target <= max_row:
        ws.delete_rows(target, max_row - target + 1)

    wb.save(destination)
    return result
