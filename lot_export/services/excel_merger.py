"""
Fusion de plusieurs fichiers audit (un par commune) en un seul classeur Excel.

Le premier fichier sert de modèle (mise en forme, largeurs de colonnes, filtres).
Pour chaque feuille, les lignes de données des fichiers suivants sont ajoutées
sous celles du modèle, sans répéter l'en-tête. Les feuilles identiques d'un
fichier à l'autre (listes déroulantes...) ne sont conservées qu'une fois.
Les lignes ne contenant que des formules (lignes « modèle » vides) sont ignorées.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.worksheet import Worksheet

from .excel_filter import AUDIT_SHEET

HEADER_ROWS = 1


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_data_row(row: Iterable[Cell]) -> bool:
    """Une ligne de données contient au moins une valeur saisie (pas seulement des formules)."""
    return any(cell.value is not None and not _is_formula(cell.value) for cell in row)


def _trim(ws: Worksheet) -> None:
    """Supprime les lignes situées après la dernière ligne de données."""
    last = HEADER_ROWS
    for index, row in enumerate(ws.iter_rows(min_row=HEADER_ROWS + 1), start=HEADER_ROWS + 1):
        if _is_data_row(row):
            last = index
    if ws.max_row > last:
        ws.delete_rows(last + 1, ws.max_row - last)


def _values(ws: Worksheet) -> List[Tuple]:
    return [row for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)]


def _same_content(a: Worksheet, b: Worksheet) -> bool:
    # comparaison des dimensions d'abord : évite de parcourir les grandes feuilles de données
    if (a.max_row, a.max_column) != (b.max_row, b.max_column):
        return False
    return _values(a) == _values(b)


def _copy_cell(cell: Cell, target: Cell, style_cache: Dict[tuple, object]) -> None:
    value = cell.value
    if _is_formula(value):
        # recale les références relatives (ex : J57 -> J3) sur la nouvelle ligne
        value = Translator(value, origin=cell.coordinate).translate_formula(target.coordinate)
    target.value = value
    if cell.has_style:
        # les styles ne sont pas partageables entre classeurs : copie attribut par attribut,
        # une seule fois par style source (les suivants réutilisent le style déjà converti)
        key = tuple(cell._style)
        if key in style_cache:
            target._style = copy(style_cache[key])
        else:
            target.font = copy(cell.font)
            target.fill = copy(cell.fill)
            target.border = copy(cell.border)
            target.alignment = copy(cell.alignment)
            target.number_format = cell.number_format
            target.protection = copy(cell.protection)
            style_cache[key] = copy(target._style)
    if cell.hyperlink is not None:
        target.hyperlink = cell.hyperlink.target


def _append_rows(source: Worksheet, target: Worksheet, min_row: int) -> None:
    style_cache: Dict[tuple, object] = {}
    next_row = target.max_row + 1
    for row in source.iter_rows(min_row=min_row):
        if not _is_data_row(row):
            continue
        for cell in row:
            if cell.value is not None or cell.has_style:
                _copy_cell(cell, target.cell(row=next_row, column=cell.column), style_cache)
        next_row += 1


def merge_audit_files(sources: Sequence[Path], destination: Path) -> int:
    """
    Fusionne les fichiers audit `sources` dans `destination`.
    Renvoie le nombre de lignes de données de la feuille Audit du fichier fusionné.
    """
    if not sources:
        raise ValueError("aucun fichier à fusionner")

    merged = load_workbook(sources[0])
    for ws in merged.worksheets:
        _trim(ws)

    for source in sources[1:]:
        wb = load_workbook(source)
        for ws in wb.worksheets:
            if ws.title not in merged.sheetnames:
                _append_rows(ws, merged.create_sheet(ws.title), min_row=1)
                continue
            target = merged[ws.title]
            if _same_content(ws, target):
                continue  # feuille commune à tous les fichiers (listes déroulantes...)
            _append_rows(ws, target, min_row=HEADER_ROWS + 1)

    merged.save(destination)
    audit = merged[AUDIT_SHEET] if AUDIT_SHEET in merged.sheetnames else merged.worksheets[0]
    return audit.max_row - HEADER_ROWS
