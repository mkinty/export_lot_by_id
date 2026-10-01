from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from lot_export.services.excel_merger import merge_audit_files


def _make_audit(path: Path, insee: str, ids: list, template_rows: int = 3) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    ws.append(["INSEE", "Valeur", "ID erreur", "Formule"])
    for row, error_id in enumerate(ids, start=2):
        ws.append([insee, row * 10, error_id, f"=B{row}*2"])
        ws.cell(row, 3).font = Font(bold=True)
    # lignes « modèle » ne contenant que des formules, comme dans les vrais fichiers audit
    for row in range(len(ids) + 2, len(ids) + 2 + template_rows):
        ws.cell(row, 4, f"=B{row}*2")
    data = wb.create_sheet("IPE")
    data.append(["Id immeuble", "INSEE"])
    data.append([f"IMM-{insee}", insee])
    wb.create_sheet("liste").append(["OK", "NOK"])
    wb.save(path)
    return path


def _rows(path: Path, sheet: str) -> list:
    return list(load_workbook(path)[sheet].iter_rows(values_only=True))


def test_appends_data_rows_and_shifts_formulas(tmp_path: Path):
    a = _make_audit(tmp_path / "audit_74143.xlsx", "74143", ["74143_1", "74143_2"])
    b = _make_audit(tmp_path / "audit_38068.xlsx", "38068", ["38068_5"])
    out = tmp_path / "merged.xlsx"

    rows = merge_audit_files([a, b], out)

    assert rows == 3
    assert _rows(out, "Audit") == [
        ("INSEE", "Valeur", "ID erreur", "Formule"),
        ("74143", 20, "74143_1", "=B2*2"),
        ("74143", 30, "74143_2", "=B3*2"),
        ("38068", 20, "38068_5", "=B4*2"),
    ]
    assert load_workbook(out)["Audit"].cell(4, 3).font.bold


def test_data_sheets_are_concatenated_and_identical_sheets_kept_once(tmp_path: Path):
    a = _make_audit(tmp_path / "a.xlsx", "74143", ["74143_1"])
    b = _make_audit(tmp_path / "b.xlsx", "38068", ["38068_5"])
    out = tmp_path / "merged.xlsx"

    merge_audit_files([a, b], out)

    assert _rows(out, "IPE") == [("Id immeuble", "INSEE"), ("IMM-74143", "74143"), ("IMM-38068", "38068")]
    assert _rows(out, "liste") == [("OK", "NOK")]


def test_single_source_is_trimmed_of_template_rows(tmp_path: Path):
    a = _make_audit(tmp_path / "a.xlsx", "74143", ["74143_1"])
    out = tmp_path / "merged.xlsx"

    assert merge_audit_files([a], out) == 1
    assert len(_rows(out, "Audit")) == 2
