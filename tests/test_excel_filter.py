from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from lot_export.services.excel_filter import ExcelFilterError, filter_audit_rows


def _make_audit(path: Path, ids: list) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    ws.append(["INSEE", "Valeur", "ID erreur", "Formule"])
    for row, error_id in enumerate(ids, start=2):
        ws.append(["74143", row * 10, error_id, f"=B{row}*2"])
    wb.create_sheet("IPE").append(["Id immeuble"])
    wb.save(path)
    return path


def test_keeps_only_requested_rows_and_shifts_formulas(tmp_path: Path):
    source = _make_audit(tmp_path / "audit_74143.xlsx", ["74143_1", "74143_2", None, "74143_3"])
    destination = tmp_path / "out.xlsx"

    result = filter_audit_rows(source, destination, ["74143_3", "74143_1"])

    assert result.kept_ids == ["74143_3", "74143_1"]
    assert result.missing_ids == []
    ws = load_workbook(destination)["Audit"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows == [
        ("INSEE", "Valeur", "ID erreur", "Formule"),
        ("74143", 20, "74143_1", "=B2*2"),
        ("74143", 50, "74143_3", "=B3*2"),
    ]


def test_other_sheets_are_preserved(tmp_path: Path):
    source = _make_audit(tmp_path / "audit.xlsx", ["74143_1"])
    destination = tmp_path / "out.xlsx"

    filter_audit_rows(source, destination, ["74143_1"])

    assert load_workbook(destination).sheetnames == ["Audit", "IPE"]


def test_reports_missing_ids_and_writes_nothing_when_no_match(tmp_path: Path):
    source = _make_audit(tmp_path / "audit.xlsx", ["74143_1"])
    destination = tmp_path / "out.xlsx"

    result = filter_audit_rows(source, destination, ["74143_9"])

    assert result.kept_ids == []
    assert result.missing_ids == ["74143_9"]
    assert not destination.exists()


def test_raises_when_column_missing(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "Audit"
    wb.active.append(["INSEE"])
    wb.save(tmp_path / "audit.xlsx")

    with pytest.raises(ExcelFilterError):
        filter_audit_rows(tmp_path / "audit.xlsx", tmp_path / "out.xlsx", ["74143_1"])
